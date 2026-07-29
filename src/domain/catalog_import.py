"""Parse a spreadsheet of services into validated drafts.

Import is deliberately a *parse and validate* step only: it never decides that a row
may be served.  Rows land as unpublished candidates so the human approval gate that
the rest of the catalog depends on stays in the loop, no matter what a spreadsheet
claims.
"""

from __future__ import annotations

import csv
import io
import unicodedata
from dataclasses import dataclass
from typing import Any

from pydantic import ValidationError

from domain.catalog_admin import ServiceDraft

MAX_IMPORT_ROWS = 500

# Column headers accepted for each draft field. Vietnamese team spreadsheets and
# English exports both appear in practice, so both are matched after normalization.
_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "name": ("name", "ten", "ten dich vu", "dich vu", "service", "service name"),
    "provider": ("provider", "nha cung cap", "don vi", "to chuc cung cap"),
    "service_type": ("service type", "servicetype", "loai kenh", "loai", "kenh", "channel type"),
    "category": ("category", "danh muc", "nhom", "linh vuc"),
    "description": ("description", "mo ta", "gioi thieu"),
    "launch_url": ("launch url", "launchurl", "url", "lien ket", "link", "deeplink", "duong dan"),
    "region": ("region", "khu vuc", "dia ban", "tinh thanh"),
    "target_user": ("target user", "targetuser", "doi tuong", "nguoi dung"),
    "organization": ("organization", "to chuc", "co quan"),
    "service_priority": ("service priority", "priority", "uu tien", "do uu tien"),
    "aliases": ("aliases", "alias", "bi danh", "ten khac", "tu khoa"),
    "intents": ("intents", "intent", "nhu cau", "y dinh"),
}

TEMPLATE_HEADERS = (
    "name",
    "provider",
    "service_type",
    "category",
    "description",
    "launch_url",
    "region",
    "target_user",
    "organization",
    "service_priority",
    "aliases",
    "intents",
)

TEMPLATE_EXAMPLE = (
    "Tên dịch vụ ví dụ",
    "Đơn vị cung cấp ví dụ",
    "oa",
    "utilities",
    "Mô tả ngắn về dịch vụ này.",
    "https://zalo.me/1234567890123456789",
    "TP.HCM; toàn quốc",
    "adult",
    "",
    "50",
    "bí danh 1; bí danh 2",
    "thanh_toan_hoa_don | tôi muốn đóng tiền điện",
)


class ImportFormatError(ValueError):
    """The uploaded file could not be read as a supported spreadsheet."""


@dataclass(frozen=True, slots=True)
class ImportRow:
    """One parsed spreadsheet row: either a valid draft or a reason it is not."""

    row_number: int
    name: str
    draft: ServiceDraft | None = None
    error: str | None = None


def _normalize_header(value: object) -> str:
    text = "" if value is None else str(value)
    decomposed = unicodedata.normalize("NFKD", text.casefold().replace("đ", "d"))
    without_marks = "".join(char for char in decomposed if not unicodedata.combining(char))
    return " ".join("".join(char if char.isalnum() else " " for char in without_marks).split())


def _map_columns(header_row: list[object]) -> dict[str, int]:
    """Map draft field names to column indexes, ignoring unknown columns."""

    lookup: dict[str, str] = {}
    for field, aliases in _COLUMN_ALIASES.items():
        for alias in aliases:
            lookup[alias] = field

    mapping: dict[str, int] = {}
    for index, raw in enumerate(header_row):
        matched = lookup.get(_normalize_header(raw))
        # First occurrence wins so a duplicated column cannot silently shadow it.
        if matched is not None and matched not in mapping:
            mapping[matched] = index
    return mapping


def _cell(row: list[object], mapping: dict[str, int], field: str) -> str:
    index = mapping.get(field)
    if index is None or index >= len(row):
        return ""
    value = row[index]
    if value is None:
        return ""
    # Excel stores unformatted numbers as floats; 50.0 must not become "50.0".
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _split_list(value: str) -> tuple[str, ...]:
    parts = (part.strip() for chunk in value.split("\n") for part in chunk.split(";"))
    return tuple(dict.fromkeys(part for part in parts if part))


def _parse_intents(value: str) -> tuple[dict[str, str], ...]:
    """Read ``intent | ví dụ câu hỏi`` entries separated by ``;`` or newlines."""

    intents: list[dict[str, str]] = []
    for entry in _split_list(value):
        intent, separator, example = entry.partition("|")
        if not separator:
            raise ValueError(f"intent {entry!r} phải theo dạng 'tên_intent | câu hỏi ví dụ'")
        intent, example = intent.strip(), example.strip()
        if not intent or not example:
            raise ValueError(f"intent {entry!r} phải có cả tên intent và câu hỏi ví dụ")
        intents.append({"intent": intent, "example_query": example})
    return tuple(intents)


def _readable_validation_error(error: ValidationError) -> str:
    problems: list[str] = []
    for detail in error.errors():
        location = ".".join(str(part) for part in detail["loc"]) or "row"
        problems.append(f"{location}: {detail['msg']}")
    return "; ".join(problems)


def _row_to_import(row: list[object], mapping: dict[str, int], row_number: int) -> ImportRow:
    name = _cell(row, mapping, "name")
    payload: dict[str, Any] = {
        "name": name,
        "provider": _cell(row, mapping, "provider"),
        "service_type": _cell(row, mapping, "service_type").casefold().replace(" ", "_"),
        "category": _cell(row, mapping, "category").casefold(),
        "description": _cell(row, mapping, "description"),
        "launch_url": _cell(row, mapping, "launch_url"),
        "region": _cell(row, mapping, "region") or None,
        "target_user": _cell(row, mapping, "target_user") or None,
        "organization": _cell(row, mapping, "organization") or None,
        "aliases": _split_list(_cell(row, mapping, "aliases")),
    }

    priority = _cell(row, mapping, "service_priority")
    if priority:
        try:
            payload["service_priority"] = int(float(priority))
        except ValueError:
            return ImportRow(row_number, name, error=f"service_priority {priority!r} không phải số")

    try:
        payload["intents"] = _parse_intents(_cell(row, mapping, "intents"))
    except ValueError as exc:
        return ImportRow(row_number, name, error=str(exc))

    try:
        draft = ServiceDraft.model_validate(payload)
    except ValidationError as exc:
        return ImportRow(row_number, name, error=_readable_validation_error(exc))
    return ImportRow(row_number, name, draft=draft)


def _rows_from_csv(content: bytes) -> list[list[object]]:
    for encoding in ("utf-8-sig", "utf-16", "cp1258", "latin-1"):
        try:
            text = content.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
        # Vietnamese Excel exports frequently use ';' because of the locale's
        # decimal comma, so the delimiter is detected rather than assumed.
        sample = text[:4096]
        try:
            dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    raise ImportFormatError("Không đọc được nội dung CSV; hãy lưu file ở dạng UTF-8.")


def _rows_from_xlsx(content: bytes) -> list[list[object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ImportFormatError("Thiếu thư viện openpyxl để đọc file .xlsx.") from exc

    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,  # Read formula results, not formula text.
        )
    except Exception as exc:
        raise ImportFormatError("Không mở được file .xlsx; file có thể bị hỏng.") from exc
    try:
        sheet = workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def parse_service_spreadsheet(content: bytes, filename: str) -> list[ImportRow]:
    """Parse an uploaded ``.xlsx``/``.csv`` file into per-row results.

    Every row is reported, valid or not, so the uploader sees exactly which lines
    need fixing instead of a single opaque failure.
    """

    if not content:
        raise ImportFormatError("File rỗng.")

    lowered = filename.casefold()
    if lowered.endswith((".xlsx", ".xlsm")):
        rows = _rows_from_xlsx(content)
    elif lowered.endswith((".csv", ".txt")):
        rows = _rows_from_csv(content)
    elif lowered.endswith(".xls"):
        raise ImportFormatError(
            "Định dạng .xls cũ không được hỗ trợ; hãy lưu lại thành .xlsx hoặc .csv."
        )
    else:
        raise ImportFormatError("Chỉ hỗ trợ file .xlsx hoặc .csv.")

    non_empty = [row for row in rows if any(str(cell).strip() for cell in row if cell is not None)]
    if not non_empty:
        raise ImportFormatError("File không có dòng nào chứa dữ liệu.")

    mapping = _map_columns(non_empty[0])
    missing = [
        field
        for field in ("name", "provider", "service_type", "category", "description", "launch_url")
        if field not in mapping
    ]
    if missing:
        raise ImportFormatError(
            "Thiếu cột bắt buộc: " + ", ".join(missing) + ". Tải file mẫu để xem đúng định dạng."
        )

    data_rows = non_empty[1:]
    if len(data_rows) > MAX_IMPORT_ROWS:
        raise ImportFormatError(
            f"File có {len(data_rows)} dòng, vượt giới hạn {MAX_IMPORT_ROWS} dòng mỗi lần import."
        )

    results: list[ImportRow] = []
    seen_urls: dict[str, int] = {}
    for offset, row in enumerate(data_rows, start=2):
        parsed = _row_to_import(row, mapping, offset)
        if parsed.draft is not None:
            url = parsed.draft.launch_url
            duplicate_of = seen_urls.get(url)
            if duplicate_of is not None:
                parsed = ImportRow(
                    parsed.row_number,
                    parsed.name,
                    error=f"launch_url trùng với dòng {duplicate_of} trong cùng file",
                )
            else:
                seen_urls[url] = parsed.row_number
        results.append(parsed)
    return results


def template_rows() -> tuple[tuple[str, ...], tuple[str, ...]]:
    """Header and one example row for the downloadable import template."""

    return TEMPLATE_HEADERS, TEMPLATE_EXAMPLE
